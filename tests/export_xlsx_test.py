"""Smoke test for VotingServer.export_xlsx (Faz 3 technical workbook).

Injects a known config + votes, exports to a temp .xlsx, reopens it with openpyxl
and asserts sheet names, ranking order and a few cell values. Run:
    python tests/export_xlsx_test.py
"""
import os
import sys
import tempfile

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from backend.voting_server import VotingServer  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

log = []
def check(name, ok, extra=""):
    log.append(("PASS" if ok else "FAIL") + f"  {name}" + (f" — {extra}" if extra else ""))
    return ok

srv = VotingServer()
srv.config = {
    "scale": 10,
    "categories": [
        {"id": "overall", "label": "Genel", "weight": 50, "enabled": True},
        {"id": "fun",     "label": "Eğlence", "weight": 50, "enabled": True},
        {"id": "off",     "label": "Kapalı", "weight": 10, "enabled": False},
    ],
    "groups": {
        "jury":     {"enabled": True, "weight": 60, "label": "Jüri"},
        "audience": {"enabled": True, "weight": 40, "label": "Seyirci"},
        "team":     {"enabled": False, "weight": 10, "label": "Ekip"},
    },
}
# Game A: high jury scores; Game B: lower -> A should rank first.
srv.votes = [
    {"group": "jury",     "game_id": "A", "game_name": "Alpha", "voter": "j1", "scores": {"overall": 10, "fun": 9}, "ts": 1000},
    {"group": "jury",     "game_id": "A", "game_name": "Alpha", "voter": "j2", "scores": {"overall": 8,  "fun": 9}, "ts": 1001},
    {"group": "audience", "game_id": "A", "game_name": "Alpha", "voter": "a1", "scores": {"overall": 7,  "fun": 8}, "ts": 1002},
    {"group": "jury",     "game_id": "B", "game_name": "Beta",  "voter": "j1", "scores": {"overall": 5,  "fun": 4}, "ts": 1003},
    {"group": "audience", "game_id": "B", "game_name": "Beta",  "voter": "a1", "scores": {"overall": 6,  "fun": 5}, "ts": 1004},
]

tmp = os.path.join(tempfile.gettempdir(), "jamdeck_export_test.xlsx")
if os.path.exists(tmp):
    os.remove(tmp)

res = srv.export_xlsx(tmp)
check("export_xlsx returns ok", res.get("ok") is True, str(res))
check("file written", os.path.exists(tmp) and os.path.getsize(tmp) > 0)

wb = load_workbook(tmp)
expected = ["Özet", "Kategori Kırılımı", "Grup Kırılımı", "Ham Oylar"]
check("4 expected sheets", wb.sheetnames == expected, str(wb.sheetnames))

# --- Özet: header + ranking + per-group columns ---
ws = wb["Özet"]
hdr = [c.value for c in ws[1]]
check("Özet header starts correctly", hdr[:4] == ["#", "Oyun", "Ağırlıklı Toplam (/10)", "Toplam Oy"], str(hdr))
check("Özet has per-group avg columns", "Jüri Ort." in hdr and "Seyirci Ort." in hdr, str(hdr))
check("Ekip (disabled) excluded from groups", "Ekip Ort." not in hdr, str(hdr))
r2 = [c.value for c in ws[2]]
r3 = [c.value for c in ws[3]]
check("rank 1 is Alpha (highest weighted)", r2[1] == "Alpha", f"{r2[1]} weighted={r2[2]}")
check("rank 2 is Beta", r3[1] == "Beta", str(r3[1]))
check("Alpha weighted > Beta weighted", (r2[2] or 0) > (r3[2] or 0), f"{r2[2]} vs {r3[2]}")
check("Alpha total votes = 3", r2[3] == 3, str(r2[3]))

# --- Kategori Kırılımı: only enabled categories ---
ws = wb["Kategori Kırılımı"]
chdr = [c.value for c in ws[1]]
check("Kategori sheet excludes disabled category", all("Kapalı" not in str(h) for h in chdr), str(chdr))
check("Kategori sheet has 2 enabled cats", len(chdr) == 3, str(chdr))  # Oyun + 2 cats

# --- Grup Kırılımı: rows per (enabled group x game) ---
ws = wb["Grup Kırılımı"]
body = [[c.value for c in row] for row in ws.iter_rows(min_row=2)]
check("Grup Kırılımı has 4 rows (2 groups x 2 games)", len(body) == 4, f"rows={len(body)}")

# --- Ham Oylar: one row per raw vote ---
ws = wb["Ham Oylar"]
raw = [[c.value for c in row] for row in ws.iter_rows(min_row=2)]
check("Ham Oylar has 5 rows (all votes)", len(raw) == 5, f"rows={len(raw)}")

# --- reset_votes wipes votes (counter/results go empty) ---
rr = srv.reset_votes()
check("reset_votes returns ok", rr.get("ok") is True, str(rr))
check("votes empty after reset", srv.votes == [], f"len={len(srv.votes)}")
check("results empty after reset", srv.get_results().get("games") == [], str(srv.get_results().get("games")))

print("\n=== EXPORT XLSX TEST ===")
print("\n".join(log))
fails = sum(1 for l in log if l.startswith("FAIL"))
print(f"\n{len(log)-fails}/{len(log)} passed")
sys.exit(1 if fails else 0)
